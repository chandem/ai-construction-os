from io import BytesIO
from pathlib import Path
from typing import TypedDict


class TextChunk(TypedDict):
    content: str
    page_number: int | None


def extract_pages(filename: str, content_type: str, data: bytes) -> list[tuple[str, int | None]]:
    """Extract text while retaining page/sheet boundaries where available."""
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf" or content_type == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        return [(page.extract_text() or "", index + 1) for index, page in enumerate(reader.pages)]

    if suffix == ".docx":
        from docx import Document

        doc = Document(BytesIO(data))
        return [("\n".join(p.text for p in doc.paragraphs), None)]

    if suffix in {".xlsx", ".xls"}:
        import openpyxl

        workbook = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
        pages: list[tuple[str, int | None]] = []
        for sheet in workbook.worksheets:
            lines = [f"[Sheet: {sheet.title}]"]
            for row in sheet.iter_rows(values_only=True):
                values = [str(v) for v in row if v is not None]
                if values:
                    lines.append(" | ".join(values))
            pages.append(("\n".join(lines), None))
        return pages

    if suffix in {".txt", ".csv"} or content_type.startswith("text/"):
        return [(data.decode("utf-8", errors="replace"), None)]

    raise ValueError("Unsupported document format")


def extract_text(filename: str, content_type: str, data: bytes) -> tuple[str, int | None]:
    pages = extract_pages(filename, content_type, data)
    return "\n\n".join(text for text, _ in pages), len(pages) if pages and any(page is not None for _, page in pages) else None


def chunk_text_with_metadata(pages: list[tuple[str, int | None]], chunk_size: int = 1800, overlap: int = 250) -> list[TextChunk]:
    """Create searchable chunks without losing the source page."""
    chunks: list[TextChunk] = []

    for page_text, page_number in pages:
        normalized = " ".join(page_text.split())
        if not normalized:
            continue

        start = 0
        while start < len(normalized):
            end = min(len(normalized), start + chunk_size)
            chunks.append({"content": normalized[start:end], "page_number": page_number})
            if end == len(normalized):
                break
            start = max(0, end - overlap)

    return chunks


def chunk_text(text: str, chunk_size: int = 1800, overlap: int = 250) -> list[str]:
    """Backward-compatible plain text chunking."""
    return [chunk["content"] for chunk in chunk_text_with_metadata([(text, None)], chunk_size, overlap)]
